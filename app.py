import streamlit as st
import pandas as pd
import plotly.express as px
import psycopg2
from zoneinfo import ZoneInfo

# Paleta de colores para las gráficas (se ve bien en tema oscuro y claro)
PALETA = [
    "#D98327", "#ECE8AF", "#4A943E", "#186049", "#287A74", "#55A9A0",
    "#AEEED3", "#9BE15D", "#FFF8B0", "#030164", "#363199", "#D6FB61",
]
MESES_ES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun",
            "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]


def etiqueta_mes(fechas):
    """'Mar 2026' para cada fecha (en español)."""
    return fechas.dt.month.map(lambda m: MESES_ES[m - 1]) + " " + fechas.dt.year.astype(str)

st.set_page_config(
    page_title="Inventario de Materiales",
    page_icon="📦",
    layout="wide",
)
# ==========================================
# LOGIN
# ==========================================
import streamlit_authenticator as stauth

credentials = {
    "usernames": {
        user: dict(data)
        for user, data in st.secrets["credentials"]["usernames"].items()
    }
}

authenticator = stauth.Authenticate(
    credentials,
    st.secrets["cookie"]["name"],
    st.secrets["cookie"]["key"],
    st.secrets["cookie"]["expiry_days"],
    auto_hash=False,
)

try:
    authenticator.login()
except Exception as e:
    st.error(e)

if st.session_state.get("authentication_status") is False:
    st.error("Usuario o contraseña incorrectos")
    st.stop()
elif st.session_state.get("authentication_status") is None:
    st.warning("Por favor ingresa tu usuario y contraseña")
    st.stop()

authenticator.logout("Cerrar sesión", "sidebar")
st.sidebar.write(f'Bienvenido, {st.session_state.get("name")}')


@st.cache_data(ttl=300)
def cargar_inventario():
    database_url = st.secrets["DATABASE_URL"]
    conexion = psycopg2.connect(database_url, sslmode="require")
    df = pd.read_sql("SELECT * FROM inventario", conexion)
    conexion.close()
    return df


@st.cache_data(ttl=300)
def cargar_salidas():
    database_url = st.secrets["DATABASE_URL"]
    conexion = psycopg2.connect(database_url, sslmode="require")
    df = pd.read_sql(
        "SELECT * FROM inv_salida",
        conexion,
        parse_dates=["fecha_movimiento"],
    )
    conexion.close()
    return df


@st.cache_data(ttl=300)
def cargar_servicios():
    """Entradas de PO de ítems que no son de inventario (servicios).
    Si la tabla todavía no existe, devuelve una tabla vacía."""
    database_url = st.secrets["DATABASE_URL"]
    conexion = psycopg2.connect(database_url, sslmode="require")
    try:
        df = pd.read_sql(
            "SELECT * FROM inv_entrada WHERE es_servicio",
            conexion,
            parse_dates=["fecha_movimiento"],
        )
    except Exception:
        df = pd.DataFrame()
    conexion.close()
    return df


@st.cache_data(ttl=300)
def obtener_ultima_actualizacion():
    database_url = st.secrets["DATABASE_URL"]
    conexion = psycopg2.connect(database_url, sslmode="require")
    cur = conexion.cursor()
    cur.execute(
        "SELECT ultima_actualizacion FROM metadata_actualizacion "
        "ORDER BY id DESC LIMIT 1;"
    )
    resultado = cur.fetchone()
    conexion.close()
    return resultado[0] if resultado else None


# ==========================================
# EXPORTAR: TABLA "CONSUMO POR MES" (cantidad y pesos por ítem y equipo)
# ==========================================

def construir_excel_consumo(df, serv=None):
    """
    Arma el libro de Excel con la tabla de consumo por mes:
    una fila por año + bodega + ítem + equipo (Rig), con cantidad y valor
    en pesos de cada mes, y hojas por equipo, por tercero y el detalle.
    """
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter as L

    d = df.copy()
    d["anio"] = d["fecha_movimiento"].dt.year
    d["mes"] = d["fecha_movimiento"].dt.month
    for c in ("tercero", "rig", "proyecto", "centro_costo", "bodega_nombre"):
        d[c] = d[c].fillna("SIN DATO")

    meses = ["Ene", "Feb", "Mar", "Abr", "May", "Jun",
             "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

    f_head = Font(bold=True, color="FFFFFF", name="Arial")
    fill_a = PatternFill("solid", fgColor="1F3864")
    fill_b = PatternFill("solid", fgColor="2F5597")
    fill_tot = PatternFill("solid", fgColor="D9E1F2")
    f_norm = Font(name="Arial", size=10)
    f_bold = Font(name="Arial", size=10, bold=True)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()

    # ---------- Hoja 1: Consumo por mes ----------
    ws = wb.active
    ws.title = "Consumo por mes"
    ws["A1"] = "Consumo por ítem, equipo (Rig), año y mes — cantidad y pesos"
    ws["A1"].font = Font(bold=True, size=13, name="Arial")

    claves = ["anio", "bodega_nombre", "item_codigo", "item_descripcion", "rig"]
    agr = (
        d.groupby(claves + ["mes"])
        .agg(cant=("cantidad", "sum"), val=("costo_total", "sum"))
        .reset_index()
    )
    ultimo = d.groupby(claves)["fecha_movimiento"].max()
    grupos = sorted(
        {tuple(r) for r in agr[claves].itertuples(index=False, name=None)},
        key=lambda t: (-t[0], t[1], t[2], t[4]),
    )
    tabla = {(*r[:5], r[5]): (r[6], r[7]) for r in agr.itertuples(index=False, name=None)}

    fijas = ["Año", "Bodega", "Item", "Descripción", "Equipo (Rig)",
             "Último movimiento", "Costo unit. prom."]
    nf = len(fijas)
    fm = nf + 1                       # primera columna de meses
    qt = fm + 24                      # total cantidad
    pt = qt + 1                       # total pesos

    for c, t in enumerate(fijas, 1):
        ws.merge_cells(start_row=3, start_column=c, end_row=4, end_column=c)
        ws.cell(3, c, t)
    for k, m in enumerate(meses):
        c = fm + 2 * k
        ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + 1)
        ws.cell(3, c, m)
        ws.cell(4, c, "Cant.")
        ws.cell(4, c + 1, "Valor $")
    ws.merge_cells(start_row=3, start_column=qt, end_row=3, end_column=pt)
    ws.cell(3, qt, "TOTAL")
    ws.cell(4, qt, "Cant.")
    ws.cell(4, pt, "Valor $")
    for c in range(1, pt + 1):
        for r in (3, 4):
            x = ws.cell(r, c)
            x.font = f_head
            x.alignment = centro
            x.fill = fill_b if (c >= fm and c < qt and ((c - fm) // 2) % 2) else fill_a

    primera = 5
    for i, g in enumerate(grupos, primera):
        anio, bod, cod, desc, rig = g
        ws.cell(i, 1, int(anio))
        ws.cell(i, 2, bod)
        ws.cell(i, 3, cod)
        ws.cell(i, 4, desc)
        ws.cell(i, 5, rig)
        ult = ultimo[g]
        ws.cell(i, 6, ult.to_pydatetime() if hasattr(ult, "to_pydatetime") else ult).number_format = "dd/mm/yyyy"
        ws.cell(i, 7, f"=IFERROR({L(pt)}{i}/{L(qt)}{i},0)").number_format = "#,##0.00"
        for k in range(12):
            cant, val = tabla.get((*g, k + 1), (0, 0))
            ws.cell(i, fm + 2 * k, float(cant)).number_format = "#,##0.##"
            ws.cell(i, fm + 2 * k + 1, float(val)).number_format = "#,##0"
        qs = ",".join(f"{L(fm + 2 * k)}{i}" for k in range(12))
        ps = ",".join(f"{L(fm + 1 + 2 * k)}{i}" for k in range(12))
        ws.cell(i, qt, f"=SUM({qs})").number_format = "#,##0.##"
        ws.cell(i, pt, f"=SUM({ps})").number_format = "#,##0"
        for c in range(1, pt + 1):
            ws.cell(i, c).font = f_bold if c >= qt else f_norm

    ultima = primera + len(grupos) - 1
    t1, t2 = ultima + 1, ultima + 2
    ws.cell(t1, 4, "TOTAL (solo filas visibles según filtro)")
    ws.cell(t2, 4, "TOTAL GENERAL (sin filtro)")
    for c in range(fm, pt + 1):
        fmt = "#,##0.##" if (c - fm) % 2 == 0 else "#,##0"
        ws.cell(t1, c, f"=SUBTOTAL(109,{L(c)}{primera}:{L(c)}{ultima})").number_format = fmt
        ws.cell(t2, c, f"=SUM({L(c)}{primera}:{L(c)}{ultima})").number_format = fmt
    for c in range(1, pt + 1):
        ws.cell(t1, c).font = f_bold
        ws.cell(t1, c).fill = fill_tot
        ws.cell(t2, c).font = Font(name="Arial", size=10, italic=True)

    for c, w in zip("ABCDEFG", [7, 28, 15, 50, 16, 14, 14]):
        ws.column_dimensions[c].width = w
    for c in range(fm, pt + 1):
        ws.column_dimensions[L(c)].width = 9 if (c - fm) % 2 == 0 and c < qt else 15
    ws.column_dimensions[L(qt)].width = 10
    ws.freeze_panes = ws.cell(primera, fm)
    ws.auto_filter.ref = f"A4:{L(pt)}{ultima}"

    # ---------- Hojas por equipo y por tercero ----------
    def hoja_resumen(titulo, col, etiqueta):
        w = wb.create_sheet(titulo)
        w["A1"] = f"Costo por {etiqueta} y por mes, en pesos"
        w["A1"].font = Font(bold=True, size=13, name="Arial")
        encabezados = ["Año", etiqueta] + meses + ["TOTAL"]
        for c, t in enumerate(encabezados, 1):
            x = w.cell(3, c, t)
            x.font = f_head
            x.fill = fill_a
            x.alignment = centro
        piv = (
            d.groupby(["anio", col, "mes"])["costo_total"].sum().reset_index()
        )
        tot = piv.groupby(["anio", col])["costo_total"].sum().sort_values(ascending=False)
        orden = sorted(tot.index, key=lambda k: (-k[0], -tot[k]))
        valores = {(r[0], r[1], r[2]): r[3] for r in piv.itertuples(index=False, name=None)}
        for i, (anio, nombre) in enumerate(orden, 4):
            w.cell(i, 1, int(anio)).font = f_norm
            w.cell(i, 2, nombre).font = f_norm
            for k in range(12):
                w.cell(i, 3 + k, float(valores.get((anio, nombre, k + 1), 0))).number_format = "#,##0"
            w.cell(i, 15, f"=SUM(C{i}:N{i})").number_format = "#,##0"
            w.cell(i, 15).font = f_bold
        u = 3 + len(orden)
        w.cell(u + 1, 2, "TOTAL (filas visibles)")
        w.cell(u + 2, 2, "TOTAL GENERAL")
        for c in range(3, 16):
            w.cell(u + 1, c, f"=SUBTOTAL(109,{L(c)}4:{L(c)}{u})").number_format = "#,##0"
            w.cell(u + 2, c, f"=SUM({L(c)}4:{L(c)}{u})").number_format = "#,##0"
        for c in range(1, 16):
            w.cell(u + 1, c).font = f_bold
            w.cell(u + 1, c).fill = fill_tot
        w.column_dimensions["A"].width = 7
        w.column_dimensions["B"].width = 42
        for c in range(3, 16):
            w.column_dimensions[L(c)].width = 14
        w.freeze_panes = "C4"
        w.auto_filter.ref = f"A3:{L(15)}{u}"

    hoja_resumen("Por Equipo", "rig", "Equipo (Rig)")
    hoja_resumen("Por Tercero", "tercero", "Tercero")

    # ---------- Detalle ----------
    wd = wb.create_sheet("Detalle")
    cols = [
        ("fecha_movimiento", "Fecha"), ("bodega_nombre", "Bodega"),
        ("numero", "N° documento"), ("item_codigo", "Código"),
        ("item_descripcion", "Descripción"), ("cantidad", "Cantidad"),
        ("costo_unitario", "Costo unitario"), ("costo_total", "Costo total"),
        ("tercero", "Tercero"), ("rig", "Equipo (Rig)"),
        ("proyecto", "Proyecto"), ("centro_costo", "Centro de costo"),
    ]
    for c, (_, t) in enumerate(cols, 1):
        x = wd.cell(1, c, t)
        x.font = f_head
        x.fill = fill_a
        x.alignment = centro
    det = d.sort_values("fecha_movimiento", ascending=False)
    for i, fila in enumerate(det[[c for c, _ in cols]].itertuples(index=False, name=None), 2):
        for c, v in enumerate(fila, 1):
            if c == 1:
                v = v.to_pydatetime()
            elif v is not None and str(v) == "nan":
                v = None
            x = wd.cell(i, c, v)
            x.font = f_norm
            if c == 1:
                x.number_format = "dd/mm/yyyy"
            elif c in (7, 8):
                x.number_format = "#,##0"
    for c, w_ in zip("ABCDEFGHIJKL", [12, 28, 14, 15, 50, 10, 14, 14, 36, 16, 30, 20]):
        wd.column_dimensions[c].width = w_
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:L{len(det) + 1}"

    # ---------- Servicios (entradas de PO de ítems que no son de inventario) ----------
    if serv is not None and not serv.empty:
        sv = serv.copy()
        sv["anio"] = sv["fecha_movimiento"].dt.year
        sv["mes"] = sv["fecha_movimiento"].dt.month
        for c in ("tercero", "rig", "bodega_nombre"):
            sv[c] = sv[c].fillna("SIN DATO")

        wsv = wb.create_sheet("Servicios por mes", 1)
        wsv["A1"] = "Servicios por ítem, equipo (Rig), año y mes — valor en pesos"
        wsv["A1"].font = Font(bold=True, size=13, name="Arial")
        enc = ["Año", "Bodega", "Item", "Descripción", "Equipo (Rig)"] + meses + ["TOTAL $"]
        for c, t in enumerate(enc, 1):
            x = wsv.cell(3, c, t)
            x.font = f_head
            x.fill = fill_a
            x.alignment = centro
        kv = ["anio", "bodega_nombre", "item_codigo", "item_descripcion", "rig"]
        pv = sv.groupby(kv + ["mes"])["costo_total"].sum().reset_index()
        tv = {(*r[:5], r[5]): r[6] for r in pv.itertuples(index=False, name=None)}
        gv = sorted({tuple(r) for r in pv[kv].itertuples(index=False, name=None)},
                    key=lambda t: (-t[0], t[1], t[2], t[4]))
        for i, g in enumerate(gv, 4):
            for c, v in enumerate(g, 1):
                wsv.cell(i, c, int(v) if c == 1 else v).font = f_norm
            for k in range(12):
                wsv.cell(i, 6 + k, float(tv.get((*g, k + 1), 0))).number_format = "#,##0"
            wsv.cell(i, 18, f"=SUM(F{i}:Q{i})").number_format = "#,##0"
            wsv.cell(i, 18).font = f_bold
        ult_v = 3 + len(gv)
        wsv.cell(ult_v + 1, 4, "TOTAL (solo filas visibles según filtro)")
        wsv.cell(ult_v + 2, 4, "TOTAL GENERAL (sin filtro)")
        for c in range(6, 19):
            wsv.cell(ult_v + 1, c, f"=SUBTOTAL(109,{L(c)}4:{L(c)}{ult_v})").number_format = "#,##0"
            wsv.cell(ult_v + 2, c, f"=SUM({L(c)}4:{L(c)}{ult_v})").number_format = "#,##0"
        for c in range(1, 19):
            wsv.cell(ult_v + 1, c).font = f_bold
            wsv.cell(ult_v + 1, c).fill = fill_tot
        for c, w_ in zip("ABCDE", [7, 28, 15, 50, 16]):
            wsv.column_dimensions[c].width = w_
        for c in range(6, 19):
            wsv.column_dimensions[L(c)].width = 14
        wsv.freeze_panes = "F4"
        wsv.auto_filter.ref = f"A3:R{ult_v}"

        # ---------- Gasto total por equipo: materiales + servicios ----------
        wt = wb.create_sheet("Total por Equipo", 3)
        wt["A1"] = "Gasto por equipo (Rig) y mes, en pesos — materiales + servicios"
        wt["A1"].font = Font(bold=True, size=13, name="Arial")
        enc = ["Año", "Equipo (Rig)", "Clase"] + meses + ["TOTAL"]
        for c, t in enumerate(enc, 1):
            x = wt.cell(3, c, t)
            x.font = f_head
            x.fill = fill_a
            x.alignment = centro
        m_ = d[["anio", "rig", "mes", "costo_total"]].assign(clase="Material")
        s_ = sv[["anio", "rig", "mes", "costo_total"]].assign(clase="Servicio")
        ambos = pd.concat([m_, s_], ignore_index=True)
        pt_ = ambos.groupby(["anio", "rig", "clase", "mes"])["costo_total"].sum()
        tot_ = ambos.groupby(["anio", "rig", "clase"])["costo_total"].sum()
        orden_ = sorted(tot_.index, key=lambda k: (-k[0], k[1], k[2]))
        for i, (a_, r_, c_) in enumerate(orden_, 4):
            wt.cell(i, 1, int(a_)).font = f_norm
            wt.cell(i, 2, r_).font = f_norm
            wt.cell(i, 3, c_).font = f_norm
            for k in range(12):
                wt.cell(i, 4 + k, float(pt_.get((a_, r_, c_, k + 1), 0))).number_format = "#,##0"
            wt.cell(i, 16, f"=SUM(D{i}:O{i})").number_format = "#,##0"
            wt.cell(i, 16).font = f_bold
        ult_t = 3 + len(orden_)
        wt.cell(ult_t + 1, 3, "TOTAL (filas visibles)")
        wt.cell(ult_t + 2, 3, "TOTAL GENERAL")
        for c in range(4, 17):
            wt.cell(ult_t + 1, c, f"=SUBTOTAL(109,{L(c)}4:{L(c)}{ult_t})").number_format = "#,##0"
            wt.cell(ult_t + 2, c, f"=SUM({L(c)}4:{L(c)}{ult_t})").number_format = "#,##0"
        for c in range(1, 17):
            wt.cell(ult_t + 1, c).font = f_bold
            wt.cell(ult_t + 1, c).fill = fill_tot
        for c, w_ in zip("ABC", [7, 30, 12]):
            wt.column_dimensions[c].width = w_
        for c in range(4, 17):
            wt.column_dimensions[L(c)].width = 14
        wt.freeze_panes = "D4"
        wt.auto_filter.ref = f"A3:P{ult_t}"

        # ---------- Detalle de servicios ----------
        wds = wb.create_sheet("Detalle servicios")
        cs = [("fecha_movimiento", "Fecha"), ("bodega_nombre", "Bodega"), ("numero", "N° entrada"),
              ("po", "PO"), ("item_codigo", "Código"), ("item_descripcion", "Descripción"),
              ("detalle", "Detalle"), ("cantidad", "Cantidad"), ("costo_unitario", "Costo unitario"),
              ("costo_total", "Costo total"), ("tercero", "Proveedor"), ("rig", "Equipo (Rig)"),
              ("proyecto", "Proyecto"), ("centro_costo", "Centro de costo")]
        for c, (_, t) in enumerate(cs, 1):
            x = wds.cell(1, c, t)
            x.font = f_head
            x.fill = fill_a
            x.alignment = centro
        dv = sv.sort_values("fecha_movimiento", ascending=False)
        for i, fila in enumerate(dv[[c for c, _ in cs]].itertuples(index=False, name=None), 2):
            for c, v in enumerate(fila, 1):
                if c == 1:
                    v = v.to_pydatetime()
                elif v is not None and str(v) in ("nan", "None", "NaT"):
                    v = None
                x = wds.cell(i, c, v)
                x.font = f_norm
                if c == 1:
                    x.number_format = "dd/mm/yyyy"
                elif c in (9, 10):
                    x.number_format = "#,##0"
        for c, w_ in zip("ABCDEFGHIJKLMN", [12, 28, 14, 12, 15, 45, 40, 10, 14, 14, 36, 16, 30, 20]):
            wds.column_dimensions[c].width = w_
        wds.freeze_panes = "A2"
        wds.auto_filter.ref = f"A1:N{len(dv) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@st.cache_data
def excel_consumo_mensual(dataframe, servicios=None):
    return construir_excel_consumo(dataframe, servicios)


df_inv = cargar_inventario()
df_sal = cargar_salidas()
df_serv = cargar_servicios()
# Nombre de la bodega para mostrarlo en el historial
nombres_bodega = (
    df_inv.dropna(subset=["bodega_codigo"])
    .drop_duplicates("bodega_codigo")
    .set_index("bodega_codigo")["bodega_nombre"]
    .to_dict()
)
df_sal["bodega_nombre"] = df_sal["bodega_codigo"].map(nombres_bodega)
ultima_actualizacion = obtener_ultima_actualizacion()

tab_inv, tab_sal = st.tabs(["📦 Inventario", "📤 Salidas por tercero / equipo"])

with tab_inv:
    # ==========================================
    # FILTROS
    # ==========================================

    st.sidebar.title("📦 Filtros de inventario")
    st.sidebar.caption("Aplican a la pestaña Inventario.")

    bodegas = sorted(df_inv["bodega_nombre"].dropna().unique())
    bodega_sel = st.sidebar.multiselect("Bodega", bodegas)

    alertas = sorted(df_inv["alerta_stock"].dropna().unique())
    alerta_sel = st.sidebar.multiselect("Alerta", alertas)

    item_busqueda = st.sidebar.text_input("Buscar ítem (código o descripción)")

    df_filtrado = df_inv.copy()
    if bodega_sel:
        df_filtrado = df_filtrado[df_filtrado["bodega_nombre"].isin(bodega_sel)]
    if alerta_sel:
        df_filtrado = df_filtrado[df_filtrado["alerta_stock"].isin(alerta_sel)]
    if item_busqueda:
        mask = (
            df_filtrado["item_codigo"].astype(str).str.contains(item_busqueda, case=False, na=False)
            | df_filtrado["item_descripcion"].str.contains(item_busqueda, case=False, na=False)
        )
        df_filtrado = df_filtrado[mask]

    # ==========================================
    # ENCABEZADO Y KPIs
    # ==========================================

    st.title("Inventario de Materiales")
    st.caption("Stock, consumo promedio, alertas de reposición y valorización por bodega.")

    if ultima_actualizacion:
        # Neon guarda la hora en UTC; la convertimos a hora de Colombia
        # solo para mostrarla (el dato guardado sigue siendo UTC).
        hora_colombia = ultima_actualizacion.replace(
            tzinfo=ZoneInfo("UTC")
        ).astimezone(ZoneInfo("America/Bogota"))

        st.info(
            f"🕒 Última actualización de datos: "
            f"{hora_colombia.strftime('%d/%m/%Y %I:%M %p')}"
        )
    else:
        st.warning("🕒 Aún no hay registro de la última actualización de datos.")

    col1, col2, col3, col4, col5 = st.columns(5)

    def formato_compacto(valor):
        """Abrevia números grandes: 2,915,123 -> 2.92M, 850,000 -> 850K."""
        if valor >= 1_000_000:
            return f"${valor / 1_000_000:,.2f}M"
        if valor >= 1_000:
            return f"${valor / 1_000:,.0f}K"
        return f"${valor:,.0f}"


    col1.metric("Stock total (unidades)", f"{df_filtrado['existencias'].sum():,.0f}")
    col2.metric(
        "Valor total inventario",
        formato_compacto(df_filtrado['valor_total'].sum()),
        help=f"Valor exacto: ${df_filtrado['valor_total'].sum():,.0f}",
    )
    col3.metric(
        "Ítems críticos",
        f"{(df_filtrado['alerta_stock'] == '🔴 CRÍTICO').sum():,}",
    )
    col4.metric(
        "En reposición",
        f"{(df_filtrado['alerta_stock'] == '🟡 REPOSICIÓN').sum():,}",
    )
    col5.metric(
        "Sin rotación",
        f"{(df_filtrado['alerta_stock'] == '⚪ SIN ROTACIÓN').sum():,}",
    )

    st.divider()

    # ==========================================
    # GRÁFICOS
    # ==========================================

    col_izq, col_der = st.columns([1, 1.3])

    with col_izq:
        st.subheader("Distribución de alertas")
        conteo_alerta = df_filtrado["alerta_stock"].value_counts().reset_index()
        conteo_alerta.columns = ["Alerta", "Cantidad"]
        fig = px.pie(
            conteo_alerta, names="Alerta", values="Cantidad", hole=0.5,
            color="Alerta",
            color_discrete_map={
                "🔴 CRÍTICO": "#e63946",
                "🟡 REPOSICIÓN": "#f4a300",
                "🟢 OK": "#1a7a3a",
                "⚪ SIN ROTACIÓN": "#9ca3af",
                "🟣 ITEM SIN REGISTRO Y SIN MOVIMIENTO": "#8b5cf6",
            },
        )
        fig.update_traces(textinfo="percent+value")
        st.plotly_chart(fig, use_container_width=True)

    with col_der:
        st.subheader("Top 15 ítems por valor en inventario")
        top_valor = df_filtrado.nlargest(15, "valor_total")[["item_descripcion", "valor_total"]]
        fig2 = px.bar(
            top_valor, x="valor_total", y="item_descripcion", orientation="h",
            color="item_descripcion", color_discrete_sequence=PALETA,
        )
        fig2.update_layout(yaxis={"categoryorder": "total ascending"}, yaxis_title="",
                           xaxis_title="Valor total", showlegend=False)
        st.plotly_chart(fig2, use_container_width=True)

    st.divider()

    # ==========================================
    # MEJORA: días estimados hasta agotarse
    # ==========================================

    st.subheader("⏳ Estimado de días hasta agotar stock")
    st.caption("Basado en el consumo promedio mensual. Solo ítems con consumo activo.")

    con_consumo = df_filtrado[df_filtrado["consumo_promedio_mensual"] > 0].copy()
    con_consumo["dias_restantes"] = (
        con_consumo["existencias"] / (con_consumo["consumo_promedio_mensual"] / 30)
    ).round(0)

    tabla_dias = con_consumo.nsmallest(20, "dias_restantes")[
        ["item_codigo", "item_descripcion", "bodega_nombre", "existencias",
         "consumo_promedio_mensual", "dias_restantes", "alerta_stock"]
    ]

    st.dataframe(
        tabla_dias,
        use_container_width=True,
        hide_index=True,
        column_config={
            "item_codigo": "Código",
            "item_descripcion": "Descripción",
            "bodega_nombre": "Bodega",
            "existencias": "Stock actual",
            "consumo_promedio_mensual": st.column_config.NumberColumn("Consumo/mes", format="%.1f"),
            "dias_restantes": st.column_config.NumberColumn("Días restantes", format="%.0f"),
            "alerta_stock": "Alerta",
        },
    )

    st.divider()

    # ==========================================
    # TABLA COMPLETA (equivalente a la tabla del dashboard viejo)
    # ==========================================

    st.subheader("📋 Detalle completo de inventario")

    st.dataframe(
        df_filtrado[[
            "bodega_nombre", "item_codigo", "item_descripcion", "existencias",
            "consumo_promedio_mensual", "stock_minimos", "punto_reposicion",
            "stock_maximo", "cantidad_a_pedir", "valor_unidad", "valor_total",
            "alerta_stock",
        ]],
        use_container_width=True,
        hide_index=True,
        column_config={
            "bodega_nombre": "Bodega",
            "item_codigo": "Código",
            "item_descripcion": "Descripción",
            "existencias": "Stock actual",
            "consumo_promedio_mensual": st.column_config.NumberColumn("Consumo/mes", format="%.1f"),
            "stock_minimos": st.column_config.NumberColumn("Stock mínimo", format="%.1f"),
            "punto_reposicion": st.column_config.NumberColumn("Pto. reposición", format="%.1f"),
            "stock_maximo": st.column_config.NumberColumn("Stock máximo", format="%.1f"),
            "cantidad_a_pedir": st.column_config.NumberColumn("Cant. a pedir", format="%.1f"),
            "valor_unidad": st.column_config.NumberColumn("Valor unidad", format="$%.0f"),
            "valor_total": st.column_config.NumberColumn("Valor total", format="$%.0f"),
            "alerta_stock": "Alerta",
        },
    )

    st.caption(f"{len(df_filtrado):,} ítems mostrados de {len(df_inv):,} totales.")

    # ==========================================
    # HISTÓRICO DE SALIDAS (por ítem)
    # ==========================================

    st.divider()
    st.subheader("📈 Histórico de salidas por ítem")
    st.caption("Fuente: consulta dinámica «Salidas no contabilizadas» de Novasoft.")

    item_hist = st.selectbox(
        "Elige un ítem para ver su historial",
        options=sorted(df_sal["item_descripcion"].dropna().unique()),
        index=None,
        placeholder="Busca un ítem...",
    )

    if item_hist:
        sal_item = df_sal[df_sal["item_descripcion"] == item_hist].sort_values("fecha_movimiento")

        por_mes = (
            sal_item.assign(mes=sal_item["fecha_movimiento"].dt.to_period("M").dt.to_timestamp())
            .groupby("mes", as_index=False)["cantidad"].sum()
        )
        por_mes["Mes"] = etiqueta_mes(por_mes["mes"])
        fig3 = px.bar(
            por_mes.sort_values("mes"), x="Mes", y="cantidad", color="Mes",
            color_discrete_sequence=PALETA,
            category_orders={"Mes": list(por_mes.sort_values("mes")["Mes"])},
            title=f"Unidades entregadas por mes — {item_hist}",
        )
        fig3.update_layout(xaxis_title="", yaxis_title="Unidades", showlegend=False)
        st.plotly_chart(fig3, use_container_width=True)

        st.dataframe(
            sal_item[[
                "fecha_movimiento", "bodega_nombre", "numero", "cantidad",
                "costo_unitario", "costo_total", "tercero", "rig", "proyecto",
            ]],
            use_container_width=True,
            hide_index=True,
            column_config={
                "fecha_movimiento": st.column_config.DateColumn("Fecha", format="DD/MM/YYYY"),
                "bodega_nombre": "Bodega",
                "numero": "N° documento",
                "cantidad": "Cantidad",
                "costo_unitario": st.column_config.NumberColumn("Costo unitario", format="$%.0f"),
                "costo_total": st.column_config.NumberColumn("Costo total", format="$%.0f"),
                "tercero": "Tercero",
                "rig": "Equipo (Rig)",
                "proyecto": "Proyecto",
            },
        )

    # ==========================================
    # EXPORTAR
    # ==========================================

    @st.cache_data
    def convertir_a_excel(dataframe):
        import io
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            dataframe.to_excel(writer, index=False, sheet_name="Inventario")
        return buffer.getvalue()

    st.download_button(
        "📊 Exportar esta vista a Excel",
        data=convertir_a_excel(df_filtrado),
        file_name="inventario_filtrado.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ==========================================
# PESTAÑA: SALIDAS POR TERCERO / EQUIPO (RIG)
# ==========================================

with tab_sal:

    st.title("Salidas por tercero y equipo (Rig)")
    st.caption(
        "Fuente: consulta dinámica «Salidas no contabilizadas» de Novasoft. "
        "Cada fila es una línea de salida con su tercero, equipo, proyecto y costo."
    )

    if df_sal.empty:
        st.warning("Todavía no hay salidas cargadas.")
        st.stop()

    d = df_sal.copy()
    d["anio"] = d["fecha_movimiento"].dt.year
    d["mes"] = d["fecha_movimiento"].dt.month
    for c in ("tercero", "rig", "proyecto", "centro_costo", "bodega_nombre"):
        d[c] = d[c].fillna("SIN DATO")

    MESES = {1: "Ene", 2: "Feb", 3: "Mar", 4: "Abr", 5: "May", 6: "Jun",
             7: "Jul", 8: "Ago", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dic"}

    st.subheader("🔎 Filtros")

    f1, f2, f3 = st.columns(3)
    anios = sorted(d["anio"].unique(), reverse=True)
    anio_sel = f1.multiselect("Año", anios, default=anios[:1], key="sal_anio")
    mes_sel = f2.multiselect(
        "Mes", list(MESES), format_func=lambda m: MESES[m], key="sal_mes"
    )
    bodega_sal = f3.multiselect(
        "Bodega", sorted(d["bodega_nombre"].unique()), key="sal_bodega"
    )

    f4, f5, f6 = st.columns(3)
    rig_sel = f4.multiselect(
        "Equipo (Rig)", sorted(d["rig"].unique()), key="sal_rig"
    )
    tercero_sel = f5.multiselect(
        "Tercero (quien recibe)", sorted(d["tercero"].unique()), key="sal_tercero"
    )
    proyecto_sel = f6.multiselect(
        "Proyecto", sorted(d["proyecto"].unique()), key="sal_proyecto"
    )

    busqueda_sal = st.text_input(
        "Buscar ítem o documento (código, descripción o número SAN)",
        key="sal_busqueda",
    )

    ds = d
    if anio_sel:
        ds = ds[ds["anio"].isin(anio_sel)]
    if mes_sel:
        ds = ds[ds["mes"].isin(mes_sel)]
    if bodega_sal:
        ds = ds[ds["bodega_nombre"].isin(bodega_sal)]
    if rig_sel:
        ds = ds[ds["rig"].isin(rig_sel)]
    if tercero_sel:
        ds = ds[ds["tercero"].isin(tercero_sel)]
    if proyecto_sel:
        ds = ds[ds["proyecto"].isin(proyecto_sel)]
    if busqueda_sal:
        m = (
            ds["item_codigo"].astype(str).str.contains(busqueda_sal, case=False, na=False)
            | ds["item_descripcion"].astype(str).str.contains(busqueda_sal, case=False, na=False)
            | ds["numero"].astype(str).str.contains(busqueda_sal, case=False, na=False)
        )
        ds = ds[m]

    if ds.empty:
        st.info("No hay salidas con esos filtros.")
        st.stop()

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Costo total", formato_compacto(ds["costo_total"].sum()),
              help=f"Valor exacto: ${ds['costo_total'].sum():,.0f}")
    k2.metric("Unidades", f"{ds['cantidad'].sum():,.0f}")
    k3.metric("Documentos", f"{ds['numero'].nunique():,}")
    k4.metric("Ítems distintos", f"{ds['item_codigo'].nunique():,}")

    st.divider()

    g1, g2 = st.columns(2)

    with g1:
        st.subheader("Costo por mes")
        por_mes = (
            ds.assign(periodo=ds["fecha_movimiento"].dt.to_period("M").dt.to_timestamp())
            .groupby("periodo", as_index=False)["costo_total"].sum()
        )
        por_mes["Mes"] = etiqueta_mes(por_mes["periodo"])
        fig_mes = px.bar(
            por_mes.sort_values("periodo"), x="Mes", y="costo_total", color="Mes",
            color_discrete_sequence=PALETA,
            category_orders={"Mes": list(por_mes.sort_values("periodo")["Mes"])},
        )
        fig_mes.update_layout(xaxis_title="", yaxis_title="Costo total", showlegend=False)
        st.plotly_chart(fig_mes, use_container_width=True)

    with g2:
        st.subheader("Costo por equipo (Rig)")
        por_rig = (
            ds.groupby("rig", as_index=False)["costo_total"].sum()
            .sort_values("costo_total", ascending=False).head(15)
        )
        fig_rig = px.bar(
            por_rig, x="costo_total", y="rig", orientation="h", color="rig",
            color_discrete_sequence=PALETA,
        )
        fig_rig.update_layout(yaxis={"categoryorder": "total ascending"},
                              yaxis_title="", xaxis_title="Costo total", showlegend=False)
        st.plotly_chart(fig_rig, use_container_width=True)

    st.subheader("Costo por tercero (top 15)")
    por_ter = (
        ds.groupby("tercero", as_index=False)["costo_total"].sum()
        .sort_values("costo_total", ascending=False).head(15)
    )
    fig_ter = px.bar(
        por_ter, x="costo_total", y="tercero", orientation="h", color="tercero",
        color_discrete_sequence=PALETA,
    )
    fig_ter.update_layout(yaxis={"categoryorder": "total ascending"},
                          yaxis_title="", xaxis_title="Costo total", showlegend=False)
    st.plotly_chart(fig_ter, use_container_width=True)

    st.divider()

    r1, r2 = st.columns(2)

    with r1:
        st.subheader("Resumen por equipo (Rig)")
        res_rig = (
            ds.groupby("rig", as_index=False)
            .agg(lineas=("item_codigo", "size"), unidades=("cantidad", "sum"),
                 costo_total=("costo_total", "sum"))
            .sort_values("costo_total", ascending=False)
        )
        st.dataframe(
            res_rig, use_container_width=True, hide_index=True,
            column_config={
                "rig": "Equipo (Rig)", "lineas": "Líneas",
                "unidades": st.column_config.NumberColumn("Unidades", format="%.0f"),
                "costo_total": st.column_config.NumberColumn("Costo total", format="$%.0f"),
            },
        )

    with r2:
        st.subheader("Resumen por tercero")
        res_ter = (
            ds.groupby("tercero", as_index=False)
            .agg(lineas=("item_codigo", "size"), unidades=("cantidad", "sum"),
                 costo_total=("costo_total", "sum"))
            .sort_values("costo_total", ascending=False)
        )
        st.dataframe(
            res_ter, use_container_width=True, hide_index=True,
            column_config={
                "tercero": "Tercero", "lineas": "Líneas",
                "unidades": st.column_config.NumberColumn("Unidades", format="%.0f"),
                "costo_total": st.column_config.NumberColumn("Costo total", format="$%.0f"),
            },
        )

    st.divider()

    st.subheader("📋 Detalle de salidas")

    detalle = ds.sort_values("fecha_movimiento", ascending=False)[[
        "fecha_movimiento", "bodega_nombre", "numero", "item_codigo",
        "item_descripcion", "cantidad", "costo_unitario", "costo_total",
        "id_tercero", "tercero", "rig", "proyecto", "centro_costo",
    ]]

    st.dataframe(
        detalle, use_container_width=True, hide_index=True,
        column_config={
            "fecha_movimiento": st.column_config.DateColumn("Fecha", format="DD/MM/YYYY"),
            "bodega_nombre": "Bodega",
            "numero": "N° documento",
            "item_codigo": "Código",
            "item_descripcion": "Descripción",
            "cantidad": st.column_config.NumberColumn("Cantidad", format="%.2f"),
            "costo_unitario": st.column_config.NumberColumn("Costo unitario", format="$%.0f"),
            "costo_total": st.column_config.NumberColumn("Costo total", format="$%.0f"),
            "id_tercero": "ID Tercero",
            "tercero": "Tercero",
            "rig": "Equipo (Rig)",
            "proyecto": "Proyecto",
            "centro_costo": "Centro de costo",
        },
    )

    st.caption(f"{len(detalle):,} líneas mostradas de {len(d):,} totales.")

    st.download_button(
        "📊 Exportar estas salidas a Excel",
        data=convertir_a_excel(detalle),
        file_name="salidas_filtradas.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="descarga_salidas",
    )

    # Servicios (entradas de PO de ítems que no son de inventario), con los
    # mismos filtros de año, mes, bodega, equipo y proyecto.
    serv_sel = df_serv.copy()
    if not serv_sel.empty:
        serv_sel["anio"] = serv_sel["fecha_movimiento"].dt.year
        serv_sel["mes"] = serv_sel["fecha_movimiento"].dt.month
        serv_sel["bodega_nombre"] = serv_sel["bodega_codigo"].map(nombres_bodega).fillna("SIN DATO")
        for c in ("rig", "proyecto", "centro_costo"):
            serv_sel[c] = serv_sel[c].fillna("SIN DATO")
        if anio_sel:
            serv_sel = serv_sel[serv_sel["anio"].isin(anio_sel)]
        if mes_sel:
            serv_sel = serv_sel[serv_sel["mes"].isin(mes_sel)]
        if bodega_sal:
            serv_sel = serv_sel[serv_sel["bodega_nombre"].isin(bodega_sal)]
        if rig_sel:
            serv_sel = serv_sel[serv_sel["rig"].isin(rig_sel)]
        if proyecto_sel:
            serv_sel = serv_sel[serv_sel["proyecto"].isin(proyecto_sel)]
        if busqueda_sal:
            m_s = (
                serv_sel["item_codigo"].astype(str).str.contains(busqueda_sal, case=False, na=False)
                | serv_sel["item_descripcion"].astype(str).str.contains(busqueda_sal, case=False, na=False)
                | serv_sel["numero"].astype(str).str.contains(busqueda_sal, case=False, na=False)
            )
            serv_sel = serv_sel[m_s]
        if tercero_sel:
            st.caption(
                "ℹ️ El filtro «Tercero» no se aplica a los servicios (allí el tercero es el proveedor)."
            )

    st.download_button(
        "📅 Exportar tabla de consumo por mes (Excel completo)",
        data=excel_consumo_mensual(ds, serv_sel if not serv_sel.empty else None),
        file_name="consumo_por_mes.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help="Una fila por ítem y equipo, con cantidad y valor en pesos de cada mes, "
             "más hojas por equipo, por tercero y el detalle. Si hay servicios cargados, "
             "agrega sus hojas y el total por equipo (materiales + servicios). "
             "Respeta los filtros de arriba.",
        key="descarga_consumo_mensual",
    )
